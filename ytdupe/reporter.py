import logging
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("ytdupe")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

PRIMARY_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
DUPLICATE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

NORMAL_FONT = Font(name="Arial", size=10)
PRIMARY_FONT = Font(name="Arial", size=10)
DUPLICATE_FONT = Font(name="Arial", size=10)

NUMBER_FORMAT = "#,##0"


def _apply_header(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _apply_normal_style(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = NORMAL_FONT
        cell.border = THIN_BORDER


def _apply_status_style(ws, row: int, max_col: int, status_col: int):
    status_val = ws.cell(row=row, column=status_col).value
    if status_val == "PRIMARY":
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = PRIMARY_FILL
            cell.font = PRIMARY_FONT
            cell.border = THIN_BORDER
    elif status_val and str(status_val).startswith("DUPLIKAT"):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = DUPLICATE_FILL
            cell.font = DUPLICATE_FONT
            cell.border = THIN_BORDER
    else:
        _apply_normal_style(ws, row, max_col)


def _format_duration(seconds_val) -> str:
    if not seconds_val or not isinstance(seconds_val, (int, float)):
        return "N/A"
    total = int(seconds_val)
    mins, secs = divmod(total, 60)
    hrs, mins = divmod(mins, 60)
    if hrs > 0:
        return f"{hrs}:{mins:02d}:{secs:02d}"
    return f"{mins}:{secs:02d}"


def _get_meta(metadata, video_id: str, field: str, default="N/A"):
    if metadata is None or metadata.empty:
        return default
    row = metadata[metadata["video_id"] == video_id]
    if row.empty:
        return default
    val = row.iloc[0].get(field, default)
    if val is None or (isinstance(val, float) and str(val) == "nan"):
        return default
    return val


def create_ringkasan_sheet(ws, total_videos, total_clusters, total_duplicate_videos, threshold, config):
    ws.title = "Ringkasan"
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 50

    headers = ["Metrik", "Nilai"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    _apply_header(ws, 1, 2)

    rows = [
        ("Total Video Dianalisis", total_videos),
        ("Jumlah Cluster Duplikat", total_clusters),
        ("Total Video Duplikat", total_duplicate_videos),
        ("Similarity Threshold", threshold),
        ("Tanggal Generate", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Channel URL", config.get("channel_url", "N/A")),
        ("Stemming Digunakan", "Ya" if config.get("use_stemming") else "Tidak"),
        ("Series Exclusion", "Ya" if config.get("exclude_series") else "Tidak"),
        ("", ""),
        ("Rekomendasi", f"Review {total_duplicate_videos} video duplikat di sheet 'Duplikat Konten'"),
    ]
    for i, (label, value) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        _apply_normal_style(ws, i, 2)


def create_duplikat_sheet(ws, clusters, metadata, video_ids, similarity_matrix):
    ws.title = "Duplikat Konten"
    headers = [
        "Group ID", "Status", "Video ID", "Judul",
        "Tanggal Publikasi", "Durasi", "Penayangan", "Similarity Score",
    ]
    widths = [10, 22, 16, 40, 18, 10, 15, 18]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        ws.cell(row=1, column=col, value=h)
        ws.column_dimensions[get_column_letter(col)].width = w
    _apply_header(ws, 1, len(headers))

    row_num = 2
    vid_to_idx = {vid: i for i, vid in enumerate(video_ids)}

    for cluster in clusters:
        for video in cluster["videos"]:
            vid = video["video_id"]
            ws.cell(row=row_num, column=1, value=cluster["group_id"])
            ws.cell(row=row_num, column=2, value=video["status"])
            ws.cell(row=row_num, column=3, value=vid)
            ws.cell(row=row_num, column=4, value=_get_meta(metadata, vid, "judul", "N/A"))
            ws.cell(row=row_num, column=5, value=_get_meta(metadata, vid, "tanggal_publikasi", "N/A"))

            dur_str = _get_meta(metadata, vid, "durasi", "N/A")
            if dur_str != "N/A":
                dur_str = _format_duration(dur_str) if isinstance(dur_str, (int, float)) else str(dur_str)
            ws.cell(row=row_num, column=6, value=dur_str)

            views = _get_meta(metadata, vid, "penayangan", 0)
            views_cell = ws.cell(row=row_num, column=7, value=int(views) if views != "N/A" else 0)
            views_cell.number_format = NUMBER_FORMAT

            sim_score = video["similarity_to_primary"]
            ws.cell(row=row_num, column=8, value=sim_score)

            _apply_status_style(ws, row_num, len(headers), status_col=2)
            row_num += 1

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_num - 1}"


def create_transkrip_sheet(ws, transcripts):
    ws.title = "Transkrip"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 80

    ws.cell(row=1, column=1, value="Video ID")
    ws.cell(row=1, column=2, value="Preview Transkrip (200 karakter)")
    _apply_header(ws, 1, 2)

    row_num = 2
    for vid in sorted(transcripts.keys()):
        text = transcripts[vid]
        preview = text[:200] + "..." if len(text) > 200 else text
        ws.cell(row=row_num, column=1, value=vid)
        ws.cell(row=row_num, column=2, value=preview)
        _apply_normal_style(ws, row_num, 2)
        row_num += 1

    ws.freeze_panes = "A2"


def create_similarity_matrix_sheet(ws, similarity_matrix, video_ids):
    ws.title = "Matriks Similarity"

    ws.cell(row=1, column=1, value="")
    for col, vid in enumerate(video_ids, 2):
        ws.cell(row=1, column=col, value=vid)
    _apply_header(ws, 1, len(video_ids) + 1)

    for i, vid in enumerate(video_ids):
        ws.cell(row=i + 2, column=1, value=vid)
        ws.cell(row=i + 2, column=1).font = Font(name="Arial", size=9, bold=True)
        for j in range(len(video_ids)):
            cell = ws.cell(row=i + 2, column=j + 2, value=round(float(similarity_matrix[i][j]), 4))
            cell.font = Font(name="Arial", size=9)
            cell.number_format = "0.0000"

    for col_idx in range(1, len(video_ids) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 12

    last_row = len(video_ids) + 1
    last_col = len(video_ids) + 1
    range_str = f"B2:{get_column_letter(last_col)}{last_row}"
    ws.conditional_formatting.add(
        range_str,
        ColorScaleRule(
            start_type="num", start_value=0, start_color="FF0000",
            mid_type="num", mid_value=0.5, mid_color="FFFF00",
            end_type="num", end_value=1.0, end_color="00FF00",
        ),
    )


def generate_report(
    output_path: str,
    transcripts: dict[str, str],
    clusters: list[dict],
    metadata,
    similarity_matrix,
    video_ids: list[str],
    config: dict,
) -> str:
    wb = Workbook()

    total_videos = len(transcripts)
    total_clusters = len(clusters)
    total_duplicate_videos = sum(c["cluster_size"] - 1 for c in clusters)

    ws_ringkasan = wb.active
    create_ringkasan_sheet(
        ws_ringkasan,
        total_videos,
        total_clusters,
        total_duplicate_videos,
        config.get("similarity_threshold", 0.75),
        config,
    )

    ws_duplikat = wb.create_sheet()
    create_duplikat_sheet(ws_duplikat, clusters, metadata, video_ids, similarity_matrix)

    ws_transkrip = wb.create_sheet()
    create_transkrip_sheet(ws_transkrip, transcripts)

    if len(video_ids) <= 100 and similarity_matrix is not None:
        ws_matrix = wb.create_sheet()
        create_similarity_matrix_sheet(ws_matrix, similarity_matrix, video_ids)

    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    wb.save(output_path)
    logger.info("Laporan disimpan: %s", output_path)
    return output_path
